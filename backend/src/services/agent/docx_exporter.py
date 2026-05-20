"""
Markdown-to-DOCX exporter for GARE analysis reports.

Converts the markdown report into a professionally styled Word document
that mirrors the look of the web UI report viewer.
"""

import re
from pathlib import Path
from typing import List, Tuple

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# -- Colour palette (matches webUI dark theme mapped to print) --
BLUE_PRIMARY = RGBColor(0x1A, 0x56, 0xA8)    # Headings
BLUE_LIGHT = RGBColor(0x2B, 0x6C, 0xB0)      # Sub-headings
GREY_DARK = RGBColor(0x24, 0x29, 0x2E)        # Code text
GREY_MED = RGBColor(0x4A, 0x4A, 0x4A)         # Secondary text
RED_ACCENT = RGBColor(0xC0, 0x39, 0x2B)       # Inline code
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
TABLE_HEADER_BG = "1A56A8"
TABLE_ALT_BG = "F0F4F8"
CODE_BG = "F6F8FA"


def markdown_to_docx(markdown_text: str, output_path: str) -> str:
    """Convert markdown report text to a styled .docx file.

    Args:
        markdown_text: The full markdown report content.
        output_path: Where to write the .docx file.

    Returns:
        The output_path string.
    """
    doc = Document()
    _setup_styles(doc)
    _setup_page(doc)

    lines = markdown_text.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]

        # --- Code block ---
        if line.strip().startswith("```"):
            lang = line.strip().lstrip("`").strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            _add_code_block(doc, "\n".join(code_lines))
            continue

        # --- Table ---
        if "|" in line and line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            _add_table(doc, table_lines)
            continue

        # --- Heading ---
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            if level == 1:
                _add_heading(doc, text, level=0)
            elif level == 2:
                _add_heading(doc, text, level=1)
            else:
                _add_heading(doc, text, level=2)
            i += 1
            continue

        # --- Horizontal rule ---
        if line.strip() in ("---", "***", "___"):
            _add_horizontal_rule(doc)
            i += 1
            continue

        # --- Unordered list item ---
        if re.match(r"^\s*[-*]\s+", line):
            list_items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                indent_match = re.match(r"^(\s*)", lines[i])
                indent = len(indent_match.group(1)) if indent_match else 0
                text = re.sub(r"^\s*[-*]\s+", "", lines[i])
                list_items.append((indent, text))
                i += 1
            for indent, text in list_items:
                level = min(indent // 2, 2)
                _add_list_item(doc, text, ordered=False, level=level)
            continue

        # --- Ordered list item ---
        if re.match(r"^\s*\d+\.\s+", line):
            list_items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                text = re.sub(r"^\s*\d+\.\s+", "", lines[i])
                list_items.append(text)
                i += 1
            for text in list_items:
                _add_list_item(doc, text, ordered=True)
            continue

        # --- Blockquote ---
        if line.strip().startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            _add_blockquote(doc, " ".join(quote_lines))
            continue

        # --- Empty line ---
        if not line.strip():
            i += 1
            continue

        # --- Normal paragraph ---
        _add_paragraph(doc, line.strip())
        i += 1

    doc.save(output_path)
    return output_path


# ------------------------------------------------------------------ #
#                        Document Setup                                #
# ------------------------------------------------------------------ #

def _setup_page(doc: Document):
    """Configure page layout."""
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def _setup_styles(doc: Document):
    """Create custom styles for the report."""
    styles = doc.styles

    # -- Title --
    title_style = styles["Title"]
    title_style.font.size = Pt(26)
    title_style.font.bold = True
    title_style.font.color.rgb = BLUE_PRIMARY
    title_style.font.name = "Calibri"
    title_style.paragraph_format.space_after = Pt(4)

    # -- Heading 1 (maps to ## in the report) --
    h1 = styles["Heading 1"]
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = BLUE_PRIMARY
    h1.font.name = "Calibri"
    h1.paragraph_format.space_before = Pt(24)
    h1.paragraph_format.space_after = Pt(8)
    _set_bottom_border(h1, color="1A56A8", size="8")

    # -- Heading 2 (maps to ### in the report) --
    h2 = styles["Heading 2"]
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = BLUE_LIGHT
    h2.font.name = "Calibri"
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)

    # -- Normal --
    normal = styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.name = "Calibri"
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15


def _set_bottom_border(style, color="1A56A8", size="8"):
    """Add a bottom border to a paragraph style."""
    pPr = style.element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="{size}" w:space="1" w:color="{color}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


# ------------------------------------------------------------------ #
#                    Element Builders                                   #
# ------------------------------------------------------------------ #

def _add_heading(doc: Document, text: str, level: int):
    """Add a heading. level 0=Title, 1=Heading1, 2=Heading2."""
    if level == 0:
        p = doc.add_paragraph(style="Title")
    elif level == 1:
        p = doc.add_paragraph(style="Heading 1")
    else:
        p = doc.add_paragraph(style="Heading 2")
    _render_inline(p, text)


def _add_paragraph(doc: Document, text: str):
    """Add a normal paragraph with inline formatting."""
    p = doc.add_paragraph(style="Normal")
    _render_inline(p, text)


def _add_list_item(doc: Document, text: str, ordered: bool = False, level: int = 0):
    """Add a list item."""
    style_name = "List Number" if ordered else "List Bullet"
    if level > 0:
        style_name += f" {min(level + 1, 3)}"
        if style_name not in [s.name for s in doc.styles]:
            style_name = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style_name)
    _render_inline(p, text)


def _add_blockquote(doc: Document, text: str):
    """Add a blockquote with left border styling."""
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(1.0)
    pPr = p._element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="18" w:space="8" w:color="AAAAAA"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    run = p.add_run(text)
    run.font.italic = True
    run.font.color.rgb = GREY_MED


def _add_code_block(doc: Document, code: str):
    """Add a code block with monospace font and shaded background."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)

    # Shaded background
    pPr = p._element.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CODE_BG}" w:val="clear"/>')
    pPr.append(shd)

    # Border
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="4" w:color="D0D7DE"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="6" w:color="D0D7DE"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="4" w:color="D0D7DE"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="6" w:color="D0D7DE"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = GREY_DARK


def _add_table(doc: Document, table_lines: List[str]):
    """Parse markdown table lines and add a styled Word table."""
    rows_raw = []
    separator_idx = None

    for idx, line in enumerate(table_lines):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.match(r"^[-:]+$", c) for c in cells):
            separator_idx = idx
            continue
        rows_raw.append(cells)

    if not rows_raw:
        return

    num_cols = max(len(r) for r in rows_raw)
    # Pad rows to same column count
    for r in rows_raw:
        while len(r) < num_cols:
            r.append("")

    table = doc.add_table(rows=len(rows_raw), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for row_idx, row_data in enumerate(rows_raw):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = ""  # Clear default
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)

            if row_idx == 0:
                # Header row: white text on blue background
                _render_inline(p, cell_text, bold_override=True)
                for run in p.runs:
                    run.font.color.rgb = WHITE
                    run.font.size = Pt(10)
                shading = parse_xml(
                    f'<w:shd {nsdecls("w")} w:fill="{TABLE_HEADER_BG}" w:val="clear"/>'
                )
                cell._element.get_or_add_tcPr().append(shading)
            else:
                _render_inline(p, cell_text)
                for run in p.runs:
                    run.font.size = Pt(10)
                # Alternate row shading
                if row_idx % 2 == 0:
                    shading = parse_xml(
                        f'<w:shd {nsdecls("w")} w:fill="{TABLE_ALT_BG}" w:val="clear"/>'
                    )
                    cell._element.get_or_add_tcPr().append(shading)

    # Add spacing after table
    doc.add_paragraph()


def _add_horizontal_rule(doc: Document):
    """Add a horizontal rule."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    pPr = p._element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


# ------------------------------------------------------------------ #
#                    Inline Formatting                                 #
# ------------------------------------------------------------------ #

def _render_inline(paragraph, text: str, bold_override: bool = False):
    """Parse inline markdown (bold, inline code) and add runs to paragraph."""
    # Tokenize: **bold**, `code`, and plain text
    tokens = _tokenize_inline(text)

    for token_type, token_text in tokens:
        run = paragraph.add_run(token_text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)

        if bold_override:
            run.font.bold = True

        if token_type == "bold":
            run.font.bold = True
        elif token_type == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RED_ACCENT
        elif token_type == "bold_code":
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RED_ACCENT


def _tokenize_inline(text: str) -> List[Tuple[str, str]]:
    """Split text into (type, content) tokens for inline formatting.

    Handles: **bold**, `code`, **`bold_code`**, and plain text.
    """
    tokens = []
    # Pattern: **`...`** | **...** | `...` | plain text
    pattern = re.compile(
        r'\*\*`([^`]+)`\*\*'   # **`bold code`**
        r'|\*\*([^*]+)\*\*'    # **bold**
        r'|`([^`]+)`'          # `code`
        r'|([^*`]+)'           # plain text
        r'|(.)'                # fallback single char
    )

    for m in pattern.finditer(text):
        if m.group(1):
            tokens.append(("bold_code", m.group(1)))
        elif m.group(2):
            tokens.append(("bold", m.group(2)))
        elif m.group(3):
            tokens.append(("code", m.group(3)))
        elif m.group(4):
            tokens.append(("plain", m.group(4)))
        elif m.group(5):
            tokens.append(("plain", m.group(5)))

    return tokens
